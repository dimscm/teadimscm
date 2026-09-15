#!/usr/bin/env python3
"""Uji tabel strata di cashback.py terhadap form monitoring resmi.

Form monitoring sudah memuat kolom cashback yang dihitung kantor. Skrip ini
menghitung ulang angka itu dari tabel strata, lalu membandingkannya baris per
baris. Kalau tabel strata salah ketik — atau bandnya berubah di bulan baru —
selisihnya langsung kelihatan di sini alih-alih diam-diam masuk ke web.

Pakai:  python3 web/tools/uji_cashback.py MONITORING_IKAT_TARGET_Q3_*.xlsx
"""

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import cashback  # noqa: E402

# Tiap sheet September: (nama sheet, baris data pertama, kolom omset total,
# kolom ACH, kolom nilai cashback, tabel strata, kolom keterangan)
SHEET = [
    ("SO TPH SEP", 7, 38, 39, 44, "TPH_SO", None),
    ("GROMIN TPH SEP", 7, 38, 39, 44, "TPH_GROMIN", None),
    ("NMAD SEP - NOV", 5, 41, 42, 47, "NMAD", None),
    ("LM 600 JAKTIM", 7, 73, 75, 79, "LM600", 81),
    ("LM 600 BEKASI", 7, 73, 75, 79, "LM600", 81),
    ("SO GALON 15L SEP", 13, 13, None, 15, "GALON_SO", None),
    ("GRM GALON 15L SEP", 13, 13, None, 15, "GALON_GROMIN", None),
]

# LM 1500+330 dipisah dua ukuran, keduanya memakai tarif dari total gabungan.
SHEET_LM1500 = [("LM 1500+330 JAKTIM", 7), ("LM 1500+330 BEKASI", 7)]


def angka(x):
    return x if isinstance(x, (int, float)) else None


def uji(nama, baris, gate):
    cocok = beda = 0
    contoh = []
    for omset, ach, tercetak, tarif in baris:
        lolos = (not gate) or (ach is not None and ach >= 1)
        hitung = round(omset * tarif) if (tarif and lolos) else 0
        if abs(hitung - (tercetak or 0)) < 1:
            cocok += 1
        else:
            beda += 1
            if len(contoh) < 3:
                contoh.append((omset, ach, tercetak, hitung))
    print(f"  {nama:28} {cocok:4} cocok / {cocok + beda:4} baris" +
          ("" if not beda else f"   <-- {beda} BEDA"))
    for omset, ach, tercetak, hitung in contoh:
        a = "-" if ach is None else f"{ach:.2f}"
        print(f"      omset={omset:9,.0f} ach={a} tercetak={tercetak:,.0f} model={hitung:,.0f}")
    return beda


def main():
    if len(sys.argv) < 2:
        sys.exit("Pakai: python3 web/tools/uji_cashback.py <form monitoring .xlsx>")
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    total_beda = 0

    for nama, awal, c_omset, c_ach, c_nilai, tabel, c_ket in SHEET:
        if nama not in wb.sheetnames:
            print(f"  {nama:28} tidak ada di workbook ini")
            continue
        baris = []
        for r in list(wb[nama].iter_rows(values_only=True))[awal:]:
            # Sheet galon memakai kolom nama outlet yang berbeda dari sheet lain.
            if not (r[4] if tabel.startswith("GALON") else r[5]):
                continue
            omset = angka(r[c_omset]) or 0
            if not omset:
                continue
            # Baris GUGUR diatur aturan turun zona, bukan tabel tarif.
            if c_ket is not None and str(r[c_ket]).strip() == "GUGUR":
                continue
            t, _ = cashback.tarif(tabel, omset)
            ach = angka(r[c_ach]) if c_ach is not None else None
            baris.append((omset, ach, angka(r[c_nilai]) or 0, t))
        total_beda += uji(nama, baris, tabel in cashback.SYARAT_ACH)

    for nama, awal in SHEET_LM1500:
        if nama not in wb.sheetnames:
            continue
        b15, b33 = [], []
        for r in list(wb[nama].iter_rows(values_only=True))[awal:]:
            if not r[5]:
                continue
            total = angka(r[94]) or 0
            if not total or str(r[103]).strip() == "GUGUR":
                continue
            t, _ = cashback.tarif("LM1500", total)
            ach = angka(r[96])
            b15.append((angka(r[87]) or 0, ach, angka(r[100]) or 0, t))
            b33.append((angka(r[93]) or 0, ach, angka(r[101]) or 0, t))
        total_beda += uji(nama + " 1500ML", b15, False)
        total_beda += uji(nama + " 330ML", b33, False)

    print()
    print("Semua cocok." if not total_beda
          else f"{total_beda} baris tidak cocok — periksa tabel strata di cashback.py.")
    return 1 if total_beda else 0


if __name__ == "__main__":
    sys.exit(main())
